"""
Uploaded document handling.

A visitor can drop in a contract, a judgment or a notice, and the app will chunk
it, embed it locally and search it alongside the corpus. Uploads live in memory
for the session only. Nothing is written to disk, which keeps a public demo free
of other people's confidential documents and keeps Render's ephemeral filesystem
irrelevant.
"""
from __future__ import annotations

import io
import logging
import re
import threading
import time
import uuid
from dataclasses import dataclass

import numpy as np

from . import config, embeddings
from .store import Chunk, UploadedDocument

log = logging.getLogger("nyaya.documents")

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}

# Paragraph and section aware splitting, same shape the corpus was built with
_SECTION_PATTERNS = [
    r"^\s*Section\s+\d+[A-Z]?[.\)]",
    r"^\s*Article\s+\d+[A-Z]?[.\)]",
    r"^\s*Clause\s+\d+",
    r"^\s*\d+\.\d+[.\s]",
    r"^\s*\d+\.\s+",
    r"^\s*\(\d+\)\s+",
    r"^\s*[IVXLC]+\.\s+",
    r"^\s*(?:CHAPTER|PART|SCHEDULE|ANNEXURE)\s+[IVXLC0-9]+",
]
_SECTION_RE = re.compile("|".join(_SECTION_PATTERNS), re.MULTILINE | re.IGNORECASE)
_SENTENCE_RE = re.compile(r"(?<=[.?!])\s+(?=[A-Z(])")

MIN_CHARS = 200
MAX_CHARS = 1800
OVERLAP_CHARS = 200


class DocumentError(ValueError):
    """A user facing problem with an uploaded file."""


# ---------------------------------------------------------------- extraction

def extract_text(filename: str, payload: bytes) -> tuple[str, int]:
    """Return (text, page_count) for a supported upload."""
    lower = filename.lower()

    if lower.endswith(".pdf"):
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise DocumentError("PDF support is not installed on this server.") from exc
        try:
            reader = PdfReader(io.BytesIO(payload))
            if reader.is_encrypted:
                try:
                    reader.decrypt("")
                except Exception as exc:
                    raise DocumentError(
                        "This PDF is password protected. Remove the password and try again."
                    ) from exc
            pages = [(page.extract_text() or "") for page in reader.pages]
        except DocumentError:
            raise
        except Exception as exc:
            raise DocumentError(f"Could not read that PDF: {exc}") from exc
        return "\n\n".join(pages), len(pages)

    if lower.endswith(".docx"):
        try:
            import docx
        except ImportError as exc:
            raise DocumentError("DOCX support is not installed on this server.") from exc
        try:
            document = docx.Document(io.BytesIO(payload))
        except Exception as exc:
            raise DocumentError(f"Could not read that DOCX file: {exc}") from exc
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n\n".join(parts), 0

    if lower.endswith((".txt", ".md")):
        for encoding in ("utf-8", "utf-16", "latin-1"):
            try:
                return payload.decode(encoding), 0
            except UnicodeDecodeError:
                continue
        raise DocumentError("Could not decode that text file.")

    raise DocumentError(
        "Unsupported file type. Upload a PDF, DOCX, TXT or MD file."
    )


# ---------------------------------------------------------------- chunking

def _split_sections(paragraph: str) -> list[str]:
    matches = list(_SECTION_RE.finditer(paragraph))
    if not matches:
        return [paragraph]
    parts: list[str] = []
    last = 0
    for match in matches:
        if match.start() > last:
            head = paragraph[last : match.start()].strip()
            if head:
                parts.append(head)
        last = match.start()
    tail = paragraph[last:].strip()
    if tail:
        parts.append(tail)
    return parts or [paragraph]


def _split_long(text: str) -> list[str]:
    if len(text) <= MAX_CHARS:
        return [text]
    parts: list[str] = []
    current = ""
    for sentence in _SENTENCE_RE.split(text):
        sentence = sentence.strip()
        if not sentence:
            continue
        if len(current) + len(sentence) + 1 <= MAX_CHARS:
            current = f"{current} {sentence}".strip()
            continue
        if current:
            parts.append(current)
            tail = current[-OVERLAP_CHARS:] if len(current) > OVERLAP_CHARS else ""
            current = f"{tail} {sentence}".strip()
        else:
            # A single sentence longer than the cap, split it hard
            for i in range(0, len(sentence), MAX_CHARS):
                parts.append(sentence[i : i + MAX_CHARS])
            current = ""
    if current:
        parts.append(current)
    return parts


def _merge_short(parts: list[str]) -> list[str]:
    merged: list[str] = []
    for part in parts:
        if merged and len(merged[-1]) < MIN_CHARS:
            merged[-1] = f"{merged[-1]}\n{part}"
        else:
            merged.append(part)
    return merged


def _section_marker(text: str) -> str:
    match = _SECTION_RE.search(text.lstrip()[:120])
    return match.group(0).strip().rstrip(".)") if match else ""


def chunk_text(text: str, *, doc_id: str, title: str) -> list[Chunk]:
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    refined: list[str] = []
    for paragraph in paragraphs:
        refined.extend(_split_sections(paragraph))
    refined = _merge_short(refined)

    chunks: list[Chunk] = []
    for paragraph in refined:
        for piece in _split_long(paragraph):
            piece = piece.strip()
            if len(piece) < 40:
                continue
            index = len(chunks)
            chunks.append(
                Chunk(
                    chunk_id=f"{doc_id}-p{index}",
                    doc_id=doc_id,
                    doc_title=title,
                    doc_type="uploaded",
                    section_marker=_section_marker(piece),
                    paragraph_index=index,
                    text=piece,
                    source_note="Uploaded by you for this session",
                    origin="upload",
                )
            )
            if len(chunks) >= config.MAX_UPLOAD_CHUNKS:
                return chunks
    return chunks


# ---------------------------------------------------------------- sessions

@dataclass
class _Entry:
    document: UploadedDocument
    touched_at: float


class UploadStore:
    """A small, bounded, in memory hold for uploaded documents."""

    def __init__(self) -> None:
        self._entries: dict[str, _Entry] = {}
        self._lock = threading.Lock()

    def _evict(self) -> None:
        now = time.time()
        stale = [
            key
            for key, entry in self._entries.items()
            if now - entry.touched_at > config.UPLOAD_SESSION_TTL_SECONDS
        ]
        for key in stale:
            self._entries.pop(key, None)

        # Hard cap on memory: drop the least recently touched documents
        while len(self._entries) > config.MAX_UPLOAD_SESSIONS:
            oldest = min(self._entries.items(), key=lambda kv: kv[1].touched_at)[0]
            self._entries.pop(oldest, None)

    def put(self, document: UploadedDocument) -> None:
        with self._lock:
            self._entries[document.doc_id] = _Entry(document, time.time())
            self._evict()

    def get(self, doc_id: str) -> UploadedDocument | None:
        with self._lock:
            entry = self._entries.get(doc_id)
            if entry is None:
                return None
            entry.touched_at = time.time()
            return entry.document

    def get_many(self, doc_ids: list[str]) -> list[UploadedDocument]:
        return [d for d in (self.get(i) for i in doc_ids) if d is not None]

    def drop(self, doc_id: str) -> bool:
        with self._lock:
            return self._entries.pop(doc_id, None) is not None

    def count(self) -> int:
        with self._lock:
            self._evict()
            return len(self._entries)


upload_store = UploadStore()


def ingest(filename: str, payload: bytes) -> UploadedDocument:
    """Extract, chunk and embed an upload. Raises DocumentError for user errors."""
    if len(payload) > config.MAX_UPLOAD_BYTES:
        limit_mb = config.MAX_UPLOAD_BYTES / 1e6
        raise DocumentError(f"That file is larger than the {limit_mb:.0f} MB limit.")
    if not payload:
        raise DocumentError("That file is empty.")

    text, n_pages = extract_text(filename, payload)
    if len(text.strip()) < 100:
        raise DocumentError(
            "No readable text was found. Scanned PDFs need to be run through OCR first."
        )

    doc_id = f"up_{uuid.uuid4().hex[:12]}"
    title = re.sub(r"\.(pdf|docx|txt|md)$", "", filename, flags=re.IGNORECASE).strip()
    title = title.replace("_", " ").replace("-", " ") or "Uploaded document"

    raw_chunks = chunk_text(text, doc_id=doc_id, title=title)
    if not raw_chunks:
        raise DocumentError("That document could not be split into readable passages.")

    truncated = len(raw_chunks) >= config.MAX_UPLOAD_CHUNKS
    vectors = embeddings.embed_texts([c.text for c in raw_chunks])

    document = UploadedDocument(
        doc_id=doc_id,
        filename=filename,
        title=title,
        chunks=raw_chunks,
        vectors=vectors,
        created_at=time.time(),
        n_pages=n_pages,
        truncated=truncated,
    )
    upload_store.put(document)
    log.info("Ingested upload %s: %d chunks", doc_id, len(raw_chunks))
    return document
